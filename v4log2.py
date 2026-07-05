from docx import Document;from docx.shared import Pt;from docx.enum.text import WD_ALIGN_PARAGRAPH as A
doc=Document();doc.styles["Normal"].font.name="Arial";doc.styles["Normal"].font.size=Pt(10)
P=lambda t:doc.add_paragraph(t)
def T(h,r):t=doc.add_table(rows=1+len(r),cols=len(h));t.style="Table Grid"
 for i,x in enumerate(h):t.rows[0].cells[i].text=str(x)
 for ri,row in enumerate(r):
  for ci,v in enumerate(row):t.rows[ri+1].cells[ci].text=str(v)
 doc.add_paragraph()
for _ in range(4):doc.add_paragraph()
p=P("CampusAsk V4 版本变更说明")
p.alignment=A.CENTER;p.runs[0].font.size=Pt(22);p.runs[0].bold=True
doc.add_paragraph()
doc.add_heading("一、新增数据库表（5张）",1)
T(["序号","表名","说明"],[
["1","knowledge_gap","知识缺口表：记录AI/联网搜索/降级兜底的问答，含urls来源链接，同一问题累加search_count"],
["2","notification","通知表：帖子被评论/回复时自动推送通知给作者，含is_read已读标记"],
["3","forum_favorite","帖子收藏表：UNIQUE(user_id,post_id)约束防重复"],
["4","forum_comment_like","评论点赞表：UNIQUE(comment_id,user_id)约束防重复，点赞同步更新forum_comment.like_count"],
["5","knowledge_history","知识版本历史表：每次编辑知识点自动记录修改前的标题/内容/分类/关键词"],
])