from docx import Document;from docx.shared import Pt;from docx.enum.text import WD_ALIGN_PARAGRAPH as A
doc=Document();doc.styles["Normal"].font.name="Arial";doc.styles["Normal"].font.size=Pt(10)
P=lambda t:doc.add_paragraph(t)
for _ in range(4):doc.add_paragraph()
p=P("CampusAsk V4 版本变更说明");p.alignment=A.CENTER;p.runs[0].font.size=Pt(22);p.runs[0].bold=True
P("项目组：CampusAsk    编制人：赵丹丹    日期：2026-07-05");doc.paragraphs[-1].alignment=A.CENTER
doc.add_paragraph()
# 11 tables here...
doc.save("D:/CampusAsk/CampusAsk/V4test.docx")
print("ok")
