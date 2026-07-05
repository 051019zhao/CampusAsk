# -*- coding: utf-8 -*-
from docx import Document;from docx.shared import Pt;from docx.enum.text import WD_ALIGN_PARAGRAPH as A
doc=Document();doc.styles[chr(78)+chr(111)+chr(114)+chr(109)+chr(97)+chr(108)].font.name=chr(65)+chr(114)+chr(105)+chr(97)+chr(108);doc.styles[chr(78)+chr(111)+chr(114)+chr(109)+chr(97)+chr(108)].font.size=Pt(11)
H1=lambda t:doc.add_heading(t,1);H2=lambda t:doc.add_heading(t,2)
B=lambda t:doc.add_paragraph(t,style=chr(76)+chr(105)+chr(115)+chr(116)+chr(32)+chr(66)+chr(117)+chr(108)+chr(108)+chr(101)+chr(116))
P=lambda t:doc.add_paragraph(t)

# Read and process markdown
f=open(chr(68)+chr(58)+chr(47)+chr(67)+chr(97)+chr(109)+chr(112)+chr(117)+chr(115)+chr(65)+chr(115)+chr(107)+chr(47)+chr(67)+chr(97)+chr(109)+chr(112)+chr(117)+chr(115)+chr(65)+chr(115)+chr(107)+chr(47)+chr(50)+chr(49)+chr(55)+chr(51)+chr(50)+chr(52)+chr(55)+chr(51)+chr(55)+chr(50)+chr(52)+chr(55)+chr(51)+chr(55)+chr(50)+chr(54)+chr(49)+chr(54)+chr(50)+chr(53)+chr(50)+chr(49)+chr(52)+chr(95)+chr(86)+chr(52)+chr(95)+chr(50)+chr(49)+chr(55)+chr(51)+chr(50)+chr(52)+chr(55)+chr(51)+chr(55)+chr(50)+chr(54)+chr(49)+chr(54)+chr(50)+chr(53)+chr(50)+chr(49)+chr(52)+chr(46)+chr(109)+chr(100),chr(114),encoding=chr(117)+chr(116)+chr(102)+chr(45)+chr(56))
data=f.read();f.close()
in_code=False
BT=chr(96);DA=chr(45);ST=chr(42);PI=chr(124);SH=chr(35)
for line in data.split(chr(10)):
 s=line.strip()
 if not s or s==DA*3:continue
 if s.startswith(BT*3):in_code=not in_code;continue
 if in_code:P(chr(32)*2+s);continue
 if s.startswith(SH+chr(32)) and not s.startswith(SH*2+chr(32)):pass
 elif s.startswith(SH*2+chr(32)) and not s.startswith(SH*3+chr(32)):H1(s[3:])
 elif s.startswith(SH*3+chr(32)):H2(s[4:])
 elif s.startswith(DA+chr(32)):B(s[2:])
 elif len(s)>0 and s[0].isdigit() and chr(46)+chr(32) in s[:4]:B(s)
 elif s.startswith(PI):pass
 else:
  t=s.replace(ST*2,chr(32)).strip()
  if t:P(t)
o=chr(68)+chr(58)+chr(47)+chr(67)+chr(97)+chr(109)+chr(112)+chr(117)+chr(115)+chr(65)+chr(115)+chr(107)+chr(47)+chr(67)+chr(97)+chr(109)+chr(112)+chr(117)+chr(115)+chr(65)+chr(115)+chr(107)+chr(47)+chr(50)+chr(49)+chr(55)+chr(51)+chr(50)+chr(52)+chr(55)+chr(51)+chr(55)+chr(50)+chr(54)+chr(49)+chr(54)+chr(50)+chr(53)+chr(50)+chr(49)+chr(52)+chr(95)+chr(86)+chr(52)+chr(95)+chr(50)+chr(49)+chr(55)+chr(51)+chr(50)+chr(52)+chr(55)+chr(51)+chr(55)+chr(50)+chr(54)+chr(49)+chr(54)+chr(50)+chr(53)+chr(50)+chr(49)+chr(52)+chr(46)+chr(100)+chr(111)+chr(99)+chr(120)
doc.save(o)
import os;print(chr(68)+chr(79)+chr(78)+chr(69),os.path.getsize(o),chr(98)+chr(121)+chr(116)+chr(101)+chr(115))